"""
Turn uploaded files (PDF / DOCX / TXT) into plain text.
"""

import io


def read_pdf(file_bytes: bytes) -> str:
    import pdfplumber
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            text_chunks.append(t)
    return "\n".join(text_chunks)


def read_docx(file_bytes: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_txt(file_bytes: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def extract_text(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            return read_pdf(file_bytes)
        elif lower.endswith(".docx"):
            return read_docx(file_bytes)
        elif lower.endswith(".txt") or lower.endswith(".md"):
            return read_txt(file_bytes)
        else:
            # best-effort fallback
            return read_txt(file_bytes)
    except Exception as e:
        return f"[ERROR extracting text from {filename}: {e}]"
